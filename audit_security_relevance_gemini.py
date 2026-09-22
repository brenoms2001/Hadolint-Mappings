import csv
import hashlib
import json
import os
import time
from collections import Counter, defaultdict
from pathlib import Path

from google import genai
from google.genai import types
from openpyxl import Workbook
from openpyxl.utils import get_column_letter


# ============================================================
# CONFIGURAÇÃO
# ============================================================

BASE_DIR = Path("data")

HADOLINT_DATASET_PATH = (
    BASE_DIR
    / "output/datasets/hadolint_rules_structured.json"
)

GEMINI_AUDIT_PATH = (
    BASE_DIR
    / "output/mappings/iec/inspection/gold_standard/"
    / "gemini_audit/full/"
    / "iec_gold_standard_gemini_audit.json"
)

OUTPUT_DIR = (
    BASE_DIR
    / "output/mappings/iec/inspection/gold_standard/"
    / "gemini_audit/"
    / "security_relevance"
)

OUTPUT_JSON = (
    OUTPUT_DIR
    / "security_relevance_gemini_audit.json"
)

OUTPUT_CSV = (
    OUTPUT_DIR
    / "security_relevance_gemini_audit.csv"
)

OUTPUT_XLSX = (
    OUTPUT_DIR
    / "security_relevance_gemini_audit.xlsx"
)

OUTPUT_MD = (
    OUTPUT_DIR
    / "security_relevance_gemini_audit.md"
)

CHECKPOINT_PATH = (
    OUTPUT_DIR
    / "security_relevance_gemini_audit_checkpoint.json"
)


# ============================================================
# GEMINI
# ============================================================

MODEL_NAME = "gemini-3.1-flash-lite"

BATCH_SIZE = 5

TEMPERATURE = 0.0

MAX_RETRIES = 5

VALIDATION_RETRIES = 2

REQUEST_DELAY = 5.0

BACKOFF_BASE = 10.0

# None = todas as regras.
MAX_RULES = None

PROMPT_VERSION = (
    "iec-gemini-security-relevance-audit-v1"
)

def get_client():

    api_key = os.getenv("GEMINI_API_KEY", "")

    if not api_key:
        raise RuntimeError(
            "Nenhuma API key do Gemini encontrada. "
            "Defina a variável de ambiente GEMINI_API_KEY."
        )

    return genai.Client(
        api_key=api_key
    )


# ============================================================
# CRITÉRIO DE CLASSIFICAÇÃO
# ============================================================

SECURITY_CRITERION = """
Classifique uma regra como relacionada à segurança quando seu
objetivo primário ou substancialmente explícito é prevenir, detectar
ou reduzir uma condição que possa comprometer:

- confidencialidade;
- integridade;
- disponibilidade;
- autenticação;
- autorização ou controle de acesso;
- privilégio mínimo;
- exposição de informações;
- execução não autorizada ou perigosa;
- injeção ou execução de comandos;
- configuração segura;
- integridade da cadeia de dependências ou software;
- segurança de rede;
- ou outra propriedade de segurança claramente identificável.

Uma regra NÃO deve ser considerada security-related apenas porque
pode ter algum efeito indireto sobre segurança.

Por exemplo, regras cujo objetivo primário é:

- estilo;
- formatação;
- sintaxe;
- portabilidade;
- manutenção;
- legibilidade;
- otimização;
- redução de tamanho;
- conveniência;
- qualidade geral do código;

devem ser classificadas como non-security quando não houver um
objetivo de segurança explícito ou substancial.

Quando houver uma fronteira genuinamente ambígua e não for possível
determinar a classificação com segurança a partir do conteúdo da
regra, use uncertain.
"""


SECURITY_CATEGORIES = [
    "authentication",
    "authorization_access_control",
    "least_privilege",
    "secrets_credentials",
    "input_validation",
    "injection_command_execution",
    "information_exposure",
    "integrity",
    "availability",
    "secure_configuration",
    "dependency_supply_chain",
    "network_security",
    "other_security",
]


# ============================================================
# SCHEMA DE RESPOSTA
# ============================================================

GEMINI_RESPONSE_SCHEMA = types.Schema(
    type=types.Type.ARRAY,
    items=types.Schema(
        type=types.Type.OBJECT,
        properties={
            "source": types.Schema(
                type=types.Type.STRING
            ),
            "source_id": types.Schema(
                type=types.Type.STRING
            ),
            "security_relevance": types.Schema(
                type=types.Type.STRING,
                enum=[
                    "security_related",
                    "non_security",
                    "uncertain",
                ],
            ),
            "security_category": types.Schema(
                type=types.Type.STRING
            ),
            "confidence": types.Schema(
                type=types.Type.INTEGER
            ),
            "security_rationale": types.Schema(
                type=types.Type.STRING
            ),
            "caveat": types.Schema(
                type=types.Type.STRING
            ),
        },
        required=[
            "source",
            "source_id",
            "security_relevance",
            "security_category",
            "confidence",
            "security_rationale",
            "caveat",
        ],
    ),
)


# ============================================================
# UTILIDADES
# ============================================================

def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()

    with path.open("rb") as f:
        for chunk in iter(
            lambda: f.read(1024 * 1024),
            b"",
        ):
            digest.update(chunk)

    return digest.hexdigest()


def load_json(path: Path):
    with path.open(
        "r",
        encoding="utf-8",
    ) as f:
        return json.load(f)


def save_json(path: Path, data):
    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with path.open(
        "w",
        encoding="utf-8",
    ) as f:
        json.dump(
            data,
            f,
            ensure_ascii=False,
            indent=2,
        )


def safe_text(value):
    if value is None:
        return ""

    if isinstance(value, (dict, list)):
        return json.dumps(
            value,
            ensure_ascii=False,
        )

    return str(value)


# ============================================================
# CARREGAMENTO DAS REGRAS
# ============================================================

def load_rules():
    data = load_json(
        HADOLINT_DATASET_PATH
    )

    if not isinstance(data, dict):
        raise ValueError(
            "O dataset de regras deveria ser um objeto/dict."
        )

    rules = []

    for source_id, rule in data.items():

        if not isinstance(rule, dict):
            raise ValueError(
                f"Regra inválida para "
                f"{source_id}: esperado objeto."
            )

        source = rule.get("source")

        if source not in {
            "hadolint",
            "shellcheck",
        }:
            raise ValueError(
                f"Regra {source_id} possui "
                f"source inválido: {source!r}"
            )

        actual_id = rule.get("id")

        if actual_id != source_id:
            raise ValueError(
                f"Inconsistência de ID: chave={source_id!r}, "
                f"campo id={actual_id!r}"
            )

        rules.append(
            {
                "source": source,
                "source_id": source_id,
                "title": rule.get("title"),
                "problematic_code": rule.get(
                    "problematic_code"
                ),
                "correct_code": rule.get(
                    "correct_code"
                ),
                "rationale": rule.get(
                    "rationale"
                ),
                "exceptions": rule.get(
                    "exceptions"
                ),
            }
        )

    rules.sort(
        key=lambda r: (
            r["source"],
            r["source_id"],
        )
    )

    return rules


# ============================================================
# PROMPT
# ============================================================

def build_prompt(batch):
    lines = []

    lines.append(
        """
You are performing a PRELIMINARY classification audit for a
research project mapping Hadolint and ShellCheck rules to
IEC 62443-3-3 security requirements.

Your task here is NOT to evaluate an IEC 62443 candidate pair.

Instead, independently classify each supplied source rule according
to whether the RULE ITSELF is primarily security-related.

This classification will later be crossed with a separate Gemini
audit of semantic associations between these rules and IEC 62443.

Human review remains the final authority.

Do not infer security relevance merely because a rule happened to
receive a high semantic similarity score against an IEC requirement.

Evaluate the actual objective and behavior of the source rule using:

- title;
- problematic code;
- correct code;
- rationale;
- exceptions.

"""
    )

    lines.append(
        SECURITY_CRITERION
    )

    lines.append(
        """
Allowed security_relevance values:

- security_related
- non_security
- uncertain

If security_relevance is security_related, provide one principal
security_category from the allowed categories.

If security_relevance is non_security, use an empty string for
security_category.

If security_relevance is uncertain, use an empty string for
security_category unless a security category is genuinely useful
for describing the ambiguity.

Allowed security categories:

- authentication
- authorization_access_control
- least_privilege
- secrets_credentials
- input_validation
- injection_command_execution
- information_exposure
- integrity
- availability
- secure_configuration
- dependency_supply_chain
- network_security
- other_security

Confidence must be an integer from 1 to 5.

Return EXACTLY one result for every supplied rule.

Do not omit rules.
Do not invent rules.
Do not change source_id values.
Do not add commentary outside the JSON response.
"""
    )

    lines.append(
        "\nRULES TO CLASSIFY:\n"
    )

    for index, rule in enumerate(batch, start=1):

        lines.append(
            f"""
--- RULE {index} ---
source: {rule["source"]}
source_id: {rule["source_id"]}
title: {safe_text(rule["title"])}

problematic_code:
{safe_text(rule["problematic_code"])}

correct_code:
{safe_text(rule["correct_code"])}

rationale:
{safe_text(rule["rationale"])}

exceptions:
{safe_text(rule["exceptions"])}
"""
        )

    return "\n".join(lines)


# ============================================================
# GEMINI
# ============================================================

def is_transient_error(exc):
    text = str(exc).lower()

    transient_markers = [
        "503",
        "unavailable",
        "service unavailable",
        "high demand",
        "temporarily unavailable",
        "internal error",
        "deadline exceeded",
        "timeout",
    ]

    return any(
        marker in text
        for marker in transient_markers
    )


def is_quota_error(exc):
    text = str(exc).lower()

    quota_markers = [
        "429",
        "quota",
        "rate limit",
        "resource exhausted",
    ]

    return any(
        marker in text
        for marker in quota_markers
    )


class GeminiBatchError(Exception):

    def __init__(
        self,
        message,
        kind="unknown",
    ):
        super().__init__(message)
        self.kind = kind


def call_gemini(batch):

    prompt = build_prompt(batch)

    last_error = None

    for attempt in range(
        1,
        MAX_RETRIES + 1,
    ):

        try:

            response = get_client().models.generate_content(
                model=MODEL_NAME,
                contents=prompt,
                config=types.GenerateContentConfig(
                    temperature=TEMPERATURE,
                    response_mime_type="application/json",
                    response_schema=GEMINI_RESPONSE_SCHEMA,
                ),
            )

            if not response.text:
                raise GeminiBatchError(
                    "Gemini retornou resposta vazia.",
                    kind="invalid_response",
                )

            try:
                parsed = json.loads(
                    response.text
                )
            except json.JSONDecodeError as exc:
                raise GeminiBatchError(
                    f"Resposta não é JSON válido: {exc}",
                    kind="invalid_response",
                ) from exc

            return parsed

        except GeminiBatchError:
            raise

        except Exception as exc:

            last_error = exc

            if is_quota_error(exc):

                raise GeminiBatchError(
                    f"Erro de quota/rate limit: {exc}",
                    kind="quota",
                ) from exc

            if not is_transient_error(exc):

                raise GeminiBatchError(
                    f"Erro não transitório: {exc}",
                    kind="fatal",
                ) from exc

            if attempt < MAX_RETRIES:

                delay = (
                    BACKOFF_BASE
                    * (2 ** (attempt - 1))
                )

                print(
                    f"      Erro transitório "
                    f"(tentativa {attempt}/{MAX_RETRIES}). "
                    f"Aguardando {delay:.1f}s..."
                )

                time.sleep(delay)

    raise GeminiBatchError(
        f"Falha após {MAX_RETRIES} tentativas: "
        f"{last_error}",
        kind="transient",
    )


# ============================================================
# VALIDAÇÃO
# ============================================================

def validate_results(
    batch,
    results,
):

    if not isinstance(
        results,
        list,
    ):
        raise GeminiBatchError(
            "Resposta deveria ser uma lista.",
            kind="invalid_response",
        )

    expected_ids = {
        (
            rule["source"],
            rule["source_id"],
        )
        for rule in batch
    }

    received_ids = []

    for result in results:

        if not isinstance(
            result,
            dict,
        ):
            raise GeminiBatchError(
                "Resultado individual não é objeto.",
                kind="invalid_response",
            )

        source = result.get(
            "source"
        )

        source_id = result.get(
            "source_id"
        )

        if not source or not source_id:
            raise GeminiBatchError(
                "Resultado sem source/source_id.",
                kind="invalid_response",
            )

        received_ids.append(
            (
                source,
                source_id,
            )
        )

        relevance = result.get(
            "security_relevance"
        )

        if relevance not in {
            "security_related",
            "non_security",
            "uncertain",
        }:
            raise GeminiBatchError(
                f"security_relevance inválido: "
                f"{relevance}",
                kind="invalid_response",
            )

        confidence = result.get(
            "confidence"
        )

        if (
            not isinstance(
                confidence,
                int,
            )
            or not 1 <= confidence <= 5
        ):
            raise GeminiBatchError(
                f"Confidence inválida: "
                f"{confidence}",
                kind="invalid_response",
            )

    received_set = set(
        received_ids
    )

    duplicates = [
        pair
        for pair, count in Counter(
            received_ids
        ).items()
        if count > 1
    ]

    if duplicates:
        raise GeminiBatchError(
            f"IDs duplicados: {duplicates}",
            kind="invalid_response",
        )

    missing = (
        expected_ids
        - received_set
    )

    extra = (
        received_set
        - expected_ids
    )

    if missing:
        raise GeminiBatchError(
            f"Resultados ausentes: "
            f"{sorted(missing)}",
            kind="invalid_response",
        )

    if extra:
        raise GeminiBatchError(
            f"Resultados extras: "
            f"{sorted(extra)}",
            kind="invalid_response",
        )

    if len(results) != len(batch):
        raise GeminiBatchError(
            "Quantidade de resultados diferente "
            "da quantidade de regras.",
            kind="invalid_response",
        )

    return True


# ============================================================
# PROCESSAMENTO ADAPTATIVO
# ============================================================

def process_batch_adaptive(
    batch,
    depth=0,
):

    indent = "  " * depth

    try:

        print(
            f"{indent}Processando "
            f"{len(batch)} regra(s)..."
        )

        for attempt in range(
            1,
            VALIDATION_RETRIES + 1,
        ):

            try:

                results = call_gemini(
                    batch
                )

                validate_results(
                    batch,
                    results,
                )

                print(
                    f"{indent}OK: "
                    f"{len(results)} resultado(s)."
                )

                return results

            except GeminiBatchError as exc:

                if exc.kind == "quota":

                    raise

                if attempt < VALIDATION_RETRIES:

                    print(
                        f"{indent}Resposta inválida "
                        f"(tentativa "
                        f"{attempt}/"
                        f"{VALIDATION_RETRIES}). "
                        f"Repetindo..."
                    )

                    time.sleep(
                        REQUEST_DELAY
                    )

                else:

                    raise

    except GeminiBatchError as exc:

        if exc.kind == "quota":
            raise

        if len(batch) == 1:

            raise RuntimeError(
                "Falha em regra individual após "
                "todas as tentativas:\n"
                f"{batch[0]['source']}:"
                f"{batch[0]['source_id']}\n"
                f"{exc}"
            ) from exc

        midpoint = len(batch) // 2

        left = batch[
            :midpoint
        ]

        right = batch[
            midpoint:
        ]

        print(
            f"{indent}Falha no batch de "
            f"{len(batch)}. "
            f"Dividindo em "
            f"{len(left)} + {len(right)}..."
        )

        left_results = (
            process_batch_adaptive(
                left,
                depth + 1,
            )
        )

        time.sleep(
            REQUEST_DELAY
        )

        right_results = (
            process_batch_adaptive(
                right,
                depth + 1,
            )
        )

        return (
            left_results
            + right_results
        )


# ============================================================
# CHECKPOINT
# ============================================================

def experiment_metadata():

    return {
        "model": MODEL_NAME,
        "batch_size": BATCH_SIZE,
        "temperature": TEMPERATURE,
        "prompt_version": PROMPT_VERSION,
        "input_rule_count": None,
        "hadolint_dataset_path": str(
            HADOLINT_DATASET_PATH
        ),
        "hadolint_dataset_hash": (
            sha256_file(
                HADOLINT_DATASET_PATH
            )
        ),
        "gemini_audit_path": str(
            GEMINI_AUDIT_PATH
        ),
        "gemini_audit_hash": (
            sha256_file(
                GEMINI_AUDIT_PATH
            )
        ),
        "criterion": SECURITY_CRITERION.strip(),
        "security_categories": SECURITY_CATEGORIES,
        "purpose": (
            "Preliminary Gemini classification of "
            "source-rule security relevance. "
            "This classification is auxiliary and "
            "does not constitute the gold standard."
        ),
        "human_review_required": True,
    }


def save_checkpoint(
    metadata,
    results,
):

    payload = {
        "metadata": metadata,
        "results": results,
    }

    save_json(
        CHECKPOINT_PATH,
        payload,
    )


def load_checkpoint():

    if not CHECKPOINT_PATH.exists():
        return None

    print(
        f"Checkpoint encontrado: "
        f"{CHECKPOINT_PATH}"
    )

    data = load_json(
        CHECKPOINT_PATH
    )

    if not isinstance(
        data,
        dict,
    ):
        raise ValueError(
            "Checkpoint inválido."
        )

    return data


# ============================================================
# CROSS-ANALYSIS
# ============================================================

def load_gemini_audit():

    data = load_json(
        GEMINI_AUDIT_PATH
    )

    results = data.get(
        "results",
        []
    )

    if not isinstance(
        results,
        list,
    ):
        raise ValueError(
            "Campo results inválido no "
            "Gemini audit."
        )

    return data, results


def build_cross_analysis(
    security_results,
    gemini_results,
):

    security_by_rule = {
        (
            result["source"],
            result["source_id"],
        ): result
        for result in security_results
    }

    grouped = defaultdict(
        list
    )

    for result in gemini_results:

        key = (
            result["source"],
            result["source_id"],
        )

        if key in security_by_rule:
            grouped[key].append(
                result
            )

    association_counts = Counter()

    rule_counts = Counter()

    source_counts = defaultdict(
        Counter
    )

    rules_with_yes = []
    rules_with_yes_maybe = []
    rules_all_no = []

    for key, pairs in grouped.items():

        security_result = (
            security_by_rule[key]
        )

        relevance = (
            security_result[
                "security_relevance"
            ]
        )

        for pair in pairs:

            verdict = pair.get(
                "llm_verdict"
            )

            association_counts[
                (
                    relevance,
                    verdict,
                )
            ] += 1

            source_counts[
                relevance
            ][verdict] += 1

        verdicts = {
            pair.get("llm_verdict")
            for pair in pairs
        }

        if "YES" in verdicts:
            rules_with_yes.append(
                {
                    "source": key[0],
                    "source_id": key[1],
                    "security_relevance": relevance,
                }
            )

        if verdicts & {
            "YES",
            "MAYBE",
        }:
            rules_with_yes_maybe.append(
                {
                    "source": key[0],
                    "source_id": key[1],
                    "security_relevance": relevance,
                }
            )

        if verdicts == {"NO"}:
            rules_all_no.append(
                {
                    "source": key[0],
                    "source_id": key[1],
                    "security_relevance": relevance,
                }
            )

        for verdict in verdicts:
            rule_counts[
                (
                    relevance,
                    verdict,
                )
            ] += 1

    return {
        "association_counts": {
            f"{relevance}__{verdict}": count
            for (
                relevance,
                verdict,
            ), count in association_counts.items()
        },
        "rule_counts": {
            f"{relevance}__{verdict}": count
            for (
                relevance,
                verdict,
            ), count in rule_counts.items()
        },
        "source_counts": {
            relevance: dict(
                counts
            )
            for relevance, counts
            in source_counts.items()
        },
        "rules_with_yes": rules_with_yes,
        "rules_with_yes_maybe": (
            rules_with_yes_maybe
        ),
        "rules_all_no": rules_all_no,
    }


# ============================================================
# CSV
# ============================================================

def save_csv(
    security_results,
    gemini_results,
):

    gemini_by_rule = defaultdict(
        list
    )

    for result in gemini_results:

        key = (
            result["source"],
            result["source_id"],
        )

        gemini_by_rule[key].append(
            result
        )

    rows = []

    for result in security_results:

        key = (
            result["source"],
            result["source_id"],
        )

        pairs = gemini_by_rule.get(
            key,
            []
        )

        verdict_counts = Counter(
            pair.get("llm_verdict")
            for pair in pairs
        )

        rows.append(
            {
                "source": result["source"],
                "source_id": result["source_id"],
                "title": result["title"],
                "security_relevance": result[
                    "security_relevance"
                ],
                "security_category": result[
                    "security_category"
                ],
                "confidence": result[
                    "confidence"
                ],
                "security_rationale": result[
                    "security_rationale"
                ],
                "caveat": result[
                    "caveat"
                ],
                "gemini_yes": verdict_counts[
                    "YES"
                ],
                "gemini_maybe": verdict_counts[
                    "MAYBE"
                ],
                "gemini_no": verdict_counts[
                    "NO"
                ],
                "gemini_yes_or_maybe": (
                    verdict_counts["YES"]
                    + verdict_counts["MAYBE"]
                ),
            }
        )

    fieldnames = list(
        rows[0].keys()
    ) if rows else []

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    with OUTPUT_CSV.open(
        "w",
        encoding="utf-8",
        newline="",
    ) as f:

        writer = csv.DictWriter(
            f,
            fieldnames=fieldnames,
        )

        writer.writeheader()

        writer.writerows(
            rows
        )


# ============================================================
# XLSX
# ============================================================

def save_xlsx(
    security_results,
    gemini_results,
    metadata,
    analysis,
):

    wb = Workbook()

    ws = wb.active
    ws.title = "Rule Review"

    headers = [
        "source",
        "source_id",
        "title",
        "security_relevance",
        "security_category",
        "confidence",
        "security_rationale",
        "caveat",
        "gemini_yes",
        "gemini_maybe",
        "gemini_no",
        "gemini_yes_or_maybe",
    ]

    ws.append(headers)

    gemini_by_rule = defaultdict(
        list
    )

    for result in gemini_results:

        key = (
            result["source"],
            result["source_id"],
        )

        gemini_by_rule[key].append(
            result
        )

    for result in security_results:

        key = (
            result["source"],
            result["source_id"],
        )

        counts = Counter(
            pair.get("llm_verdict")
            for pair in gemini_by_rule.get(
                key,
                []
            )
        )

        ws.append(
            [
                result["source"],
                result["source_id"],
                result["title"],
                result["security_relevance"],
                result["security_category"],
                result["confidence"],
                result["security_rationale"],
                result["caveat"],
                counts["YES"],
                counts["MAYBE"],
                counts["NO"],
                (
                    counts["YES"]
                    + counts["MAYBE"]
                ),
            ]
        )

    for column_index, header in enumerate(
        headers,
        start=1,
    ):
        ws.column_dimensions[
            get_column_letter(
                column_index
            )
        ].width = min(
            max(
                len(header) + 2,
                15,
            ),
            40,
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # --------------------------------------------------------
    # Summary
    # --------------------------------------------------------

    summary = wb.create_sheet(
        "Summary"
    )

    summary.append(
        ["Metric", "Value"]
    )

    summary_rows = [
        (
            "Total rules",
            len(security_results),
        ),
        (
            "Security-related",
            sum(
                1
                for r in security_results
                if r["security_relevance"]
                == "security_related"
            ),
        ),
        (
            "Non-security",
            sum(
                1
                for r in security_results
                if r["security_relevance"]
                == "non_security"
            ),
        ),
        (
            "Uncertain",
            sum(
                1
                for r in security_results
                if r["security_relevance"]
                == "uncertain"
            ),
        ),
        (
            "Gemini candidate associations",
            len(gemini_results),
        ),
    ]

    for row in summary_rows:
        summary.append(row)

    summary.append([])

    summary.append(
        [
            "Security relevance",
            "Gemini verdict",
            "Associations",
        ]
    )

    for key, value in sorted(
        analysis[
            "association_counts"
        ].items()
    ):

        relevance, verdict = key.split(
            "__",
            1,
        )

        summary.append(
            [
                relevance,
                verdict,
                value,
            ]
        )

    # --------------------------------------------------------
    # Metadata
    # --------------------------------------------------------

    metadata_ws = wb.create_sheet(
        "Metadata"
    )

    metadata_ws.append(
        ["Key", "Value"]
    )

    for key, value in metadata.items():

        if isinstance(
            value,
            (dict, list),
        ):
            value = json.dumps(
                value,
                ensure_ascii=False,
            )

        metadata_ws.append(
            [
                key,
                value,
            ]
        )

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    wb.save(
        OUTPUT_XLSX
    )


# ============================================================
# MARKDOWN
# ============================================================

def save_markdown(
    security_results,
    gemini_results,
    analysis,
):

    total = len(
        security_results
    )

    counts = Counter(
        result[
            "security_relevance"
        ]
        for result in security_results
    )

    lines = []

    lines.append(
        "# Gemini Security Relevance Audit"
    )

    lines.append("")

    lines.append(
        "## Purpose"
    )

    lines.append("")

    lines.append(
        "This is a preliminary Gemini audit of the "
        "security relevance of the 594 Hadolint/ShellCheck "
        "source rules. It is an auxiliary analysis and "
        "does not define the gold standard."
    )

    lines.append("")

    lines.append(
        "Human review remains authoritative."
    )

    lines.append("")

    lines.append(
        "## Rule classification"
    )

    lines.append("")

    lines.append(
        f"- Total rules: {total}"
    )

    lines.append(
        f"- Security-related: "
        f"{counts['security_related']}"
    )

    lines.append(
        f"- Non-security: "
        f"{counts['non_security']}"
    )

    lines.append(
        f"- Uncertain: "
        f"{counts['uncertain']}"
    )

    lines.append("")

    lines.append(
        "## Gemini candidate associations"
    )

    lines.append("")

    lines.append(
        "| Security classification | Verdict | Associations |"
    )

    lines.append(
        "|---|---:|---:|"
    )

    for key, value in sorted(
        analysis[
            "association_counts"
        ].items()
    ):

        relevance, verdict = key.split(
            "__",
            1,
        )

        lines.append(
            f"| {relevance} | "
            f"{verdict} | "
            f"{value} |"
        )

    lines.append("")

    lines.append(
        "## Interpretation"
    )

    lines.append("")

    lines.append(
        "The security classification is evaluated at the "
        "source-rule level, whereas YES/MAYBE/NO is evaluated "
        "at the rule-to-IEC-candidate level. Therefore, these "
        "two dimensions must not be treated as interchangeable."
    )

    lines.append("")

    lines.append(
        "A rule classified as security-related does not imply "
        "that any particular IEC 62443 candidate is relevant. "
        "Likewise, a non-security rule receiving a semantic "
        "candidate does not become a security rule because of "
        "that candidate."
    )

    lines.append("")

    lines.append(
        "Human review remains the final authority for the "
        "gold-standard associations."
    )

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    OUTPUT_MD.write_text(
        "\n".join(lines),
        encoding="utf-8",
    )


# ============================================================
# MAIN
# ============================================================

def main():

    print("=" * 70)
    print(
        "GEMINI SECURITY RELEVANCE AUDIT"
    )
    print("=" * 70)

    # --------------------------------------------------------
    # LOAD
    # --------------------------------------------------------

    print()
    print(
        "Carregando regras..."
    )

    rules = load_rules()

    print(
        f"  Regras encontradas: "
        f"{len(rules)}"
    )

    if MAX_RULES is not None:
        rules = rules[
            :MAX_RULES
        ]

        print(
            f"  Limitando execução para "
            f"{len(rules)} regras."
        )

    if not rules:
        raise ValueError(
            "Nenhuma regra encontrada."
        )

    # --------------------------------------------------------
    # LOAD EXISTING GEMINI AUDIT
    # --------------------------------------------------------

    print()
    print(
        "Carregando auditoria Gemini existente..."
    )

    gemini_metadata, gemini_results = (
        load_gemini_audit()
    )

    print(
        f"  Associações Gemini: "
        f"{len(gemini_results)}"
    )

    # --------------------------------------------------------
    # METADATA
    # --------------------------------------------------------

    metadata = experiment_metadata()

    metadata[
        "input_rule_count"
    ] = len(rules)

    # --------------------------------------------------------
    # CHECKPOINT
    # --------------------------------------------------------

    checkpoint = (
        load_checkpoint()
    )

    results_by_key = {}

    if checkpoint:

        checkpoint_metadata = (
            checkpoint.get(
                "metadata",
                {},
            )
        )

        if (
            checkpoint_metadata.get(
                "hadolint_dataset_hash"
            )
            != metadata[
                "hadolint_dataset_hash"
            ]
        ):
            raise RuntimeError(
                "O hash do dataset mudou "
                "desde o checkpoint."
            )

        for result in checkpoint.get(
            "results",
            [],
        ):

            key = (
                result["source"],
                result["source_id"],
            )

            results_by_key[
                key
            ] = result

        print(
            f"  Resultados recuperados: "
            f"{len(results_by_key)}"
        )

    # --------------------------------------------------------
    # PROCESSAMENTO
    # --------------------------------------------------------

    remaining = [
        rule
        for rule in rules
        if (
            rule["source"],
            rule["source_id"],
        )
        not in results_by_key
    ]

    total_batches = (
        (
            len(remaining)
            + BATCH_SIZE
            - 1
        )
        // BATCH_SIZE
    )

    print()
    print(
        f"Regras restantes: "
        f"{len(remaining)}"
    )

    print(
        f"Batch size: "
        f"{BATCH_SIZE}"
    )

    print(
        f"Batches restantes: "
        f"{total_batches}"
    )

    for batch_number, start in enumerate(
        range(
            0,
            len(remaining),
            BATCH_SIZE,
        ),
        start=1,
    ):

        batch = remaining[
            start:start + BATCH_SIZE
        ]

        print()
        print(
            "=" * 70
        )

        print(
            f"Batch {batch_number}/"
            f"{total_batches}"
        )

        print(
            "=" * 70
        )

        batch_results = (
            process_batch_adaptive(
                batch
            )
        )

        for result in batch_results:

            key = (
                result["source"],
                result["source_id"],
            )

            original_rule = next(
                rule
                for rule in batch
                if (
                    rule["source"],
                    rule["source_id"],
                ) == key
            )

            merged = {
                **original_rule,
                **result,
            }

            results_by_key[
                key
            ] = merged

        checkpoint_results = sorted(
            results_by_key.values(),
            key=lambda r: (
                r["source"],
                r["source_id"],
            ),
        )

        save_checkpoint(
            metadata,
            checkpoint_results,
        )

        print(
            f"Batch {batch_number} concluído."
        )

        print(
            f"Progresso: "
            f"{len(results_by_key)}/"
            f"{len(rules)}"
        )

        if (
            batch_number
            < total_batches
        ):
            time.sleep(
                REQUEST_DELAY
            )

    # --------------------------------------------------------
    # FINAL VALIDATION
    # --------------------------------------------------------

    security_results = sorted(
        results_by_key.values(),
        key=lambda r: (
            r["source"],
            r["source_id"],
        ),
    )

    expected_keys = {
        (
            rule["source"],
            rule["source_id"],
        )
        for rule in rules
    }

    actual_keys = {
        (
            result["source"],
            result["source_id"],
        )
        for result in security_results
    }

    missing = (
        expected_keys
        - actual_keys
    )

    extra = (
        actual_keys
        - expected_keys
    )

    if missing:
        raise RuntimeError(
            f"Regras sem classificação: "
            f"{sorted(missing)}"
        )

    if extra:
        raise RuntimeError(
            f"Classificações extras: "
            f"{sorted(extra)}"
        )

    if len(security_results) != len(
        rules
    ):
        raise RuntimeError(
            "Quantidade final de classificações "
            "não corresponde à quantidade de regras."
        )

    # --------------------------------------------------------
    # CROSS ANALYSIS
    # --------------------------------------------------------

    print()
    print(
        "Construindo análise cruzada..."
    )

    analysis = build_cross_analysis(
        security_results,
        gemini_results,
    )

    # --------------------------------------------------------
    # FINAL JSON
    # --------------------------------------------------------

    final_metadata = {
        **metadata,
        "status": "completed",
        "result_count": len(
            security_results
        ),
        "gemini_candidate_count": len(
            gemini_results
        ),
        "gemini_audit_metadata": (
            gemini_metadata
        ),
    }

    final_output = {
        "metadata": final_metadata,
        "results": security_results,
        "cross_analysis": analysis,
    }

    save_json(
        OUTPUT_JSON,
        final_output,
    )

    # --------------------------------------------------------
    # OUTPUTS
    # --------------------------------------------------------

    save_csv(
        security_results,
        gemini_results,
    )

    save_xlsx(
        security_results,
        gemini_results,
        final_metadata,
        analysis,
    )

    save_markdown(
        security_results,
        gemini_results,
        analysis,
    )

    # --------------------------------------------------------
    # SUMMARY
    # --------------------------------------------------------

    counts = Counter(
        result[
            "security_relevance"
        ]
        for result in security_results
    )

    print()
    print("=" * 70)
    print(
        "SECURITY RELEVANCE AUDIT SUMMARY"
    )
    print("=" * 70)

    print(
        f"Total rules: "
        f"{len(security_results)}"
    )

    print(
        f"Security-related: "
        f"{counts['security_related']}"
    )

    print(
        f"Non-security: "
        f"{counts['non_security']}"
    )

    print(
        f"Uncertain: "
        f"{counts['uncertain']}"
    )

    print()

    print(
        "Outputs:"
    )

    print(
        f"  JSON:  {OUTPUT_JSON}"
    )

    print(
        f"  CSV:   {OUTPUT_CSV}"
    )

    print(
        f"  XLSX:  {OUTPUT_XLSX}"
    )

    print(
        f"  MD:    {OUTPUT_MD}"
    )


if __name__ == "__main__":
    main()
