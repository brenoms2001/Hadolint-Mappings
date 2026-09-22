#!/usr/bin/env python3
"""
Global descriptive analysis of the Gemini preliminary audit.

The script reads the final Gemini audit XLSX and produces:
  - JSON: machine-readable aggregate statistics
  - Markdown: human-readable report
  - XLSX: tables for inspection

Gemini remains a preliminary auditor. Human review remains authoritative.
This script does not convert Gemini verdicts into gold-standard labels.
"""

from __future__ import annotations

import json
import math
from collections import Counter
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


BASE_DIR = Path("data")

INPUT_XLSX = (
    BASE_DIR
    / "output/mappings/iec/inspection/gold_standard/"
    / "gemini_audit/full/iec_gold_standard_gemini_audit.xlsx"
)
FALLBACK_INPUT_XLSX = Path("/mnt/data/iec_gold_standard_gemini_audit.xlsx")

OUTPUT_DIR = (
    BASE_DIR
    / "output/mappings/iec/inspection/gold_standard/"
    / "gemini_audit/global"
)
OUTPUT_JSON = OUTPUT_DIR / "gemini_global_analysis.json"
OUTPUT_MD = OUTPUT_DIR / "gemini_global_analysis.md"
OUTPUT_XLSX = OUTPUT_DIR / "gemini_global_analysis.xlsx"

VERDICTS = ["YES", "MAYBE", "NO"]
EXPECTED_TOTAL = 4134
EXPECTED_SOURCE_COUNTS = {"hadolint": 529, "shellcheck": 3605}
EXPECTED_TARGET_COUNTS = {"SR": 2214, "RE": 1920}


def clean(v: Any) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def pct(n: float, d: float) -> float:
    return round(100.0 * n / d, 4) if d else 0.0


def stats(series: pd.Series) -> dict[str, Any]:
    x = pd.to_numeric(series, errors="coerce").dropna()
    if x.empty:
        return {"n": 0, "min": None, "q25": None, "median": None,
                "mean": None, "q75": None, "max": None, "std": None}
    return {
        "n": int(x.size),
        "min": round(float(x.min()), 6),
        "q25": round(float(x.quantile(.25)), 6),
        "median": round(float(x.median()), 6),
        "mean": round(float(x.mean()), 6),
        "q75": round(float(x.quantile(.75)), 6),
        "max": round(float(x.max()), 6),
        "std": round(float(x.std(ddof=1)), 6) if len(x) > 1 else 0.0,
    }


def distribution(series: pd.Series) -> list[dict[str, Any]]:
    s = series.fillna("(blank)").astype(str)
    counts = s.value_counts()
    total = len(s)
    return [
        {"value": str(k), "count": int(v), "percentage": pct(int(v), total)}
        for k, v in counts.items()
    ]


def cross_distribution(
    df: pd.DataFrame, group_col: str, value_col: str
) -> list[dict[str, Any]]:
    rows = []
    for group, g in df.groupby(group_col, dropna=False, sort=True):
        group_label = "(blank)" if pd.isna(group) else str(group)
        counts = g[value_col].fillna("(blank)").astype(str).value_counts()
        total = len(g)
        for value, count in counts.items():
            rows.append({
                "group": group_label,
                "value": str(value),
                "count": int(count),
                "percentage_within_group": pct(int(count), total),
            })
    return rows


def rule_summary(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for (source, source_id), g in df.groupby(
        ["source", "source_id"], sort=True
    ):
        g = g.sort_values("rank")
        v = set(g["llm_verdict"].dropna())
        rows.append({
            "source": source,
            "source_id": source_id,
            "source_title": g["source_title"].dropna().iloc[0]
                if g["source_title"].notna().any() else None,
            "candidate_count": len(g),
            "yes_count": int((g["llm_verdict"] == "YES").sum()),
            "maybe_count": int((g["llm_verdict"] == "MAYBE").sum()),
            "no_count": int((g["llm_verdict"] == "NO").sum()),
            "yes_or_maybe_count": int(
                g["llm_verdict"].isin(["YES", "MAYBE"]).sum()
            ),
            "has_yes": "YES" in v,
            "has_maybe": "MAYBE" in v,
            "has_yes_or_maybe": bool(v & {"YES", "MAYBE"}),
            "all_no": v == {"NO"},
            "top1_verdict": g.iloc[0]["llm_verdict"] if len(g) else None,
            "top1_target_id": g.iloc[0]["target_id"] if len(g) else None,
            "top1_target_type": g.iloc[0]["target_type"] if len(g) else None,
        })
    return pd.DataFrame(rows)


def top_targets(df: pd.DataFrame, verdicts: set[str], limit: int = 30):
    g = df[df["llm_verdict"].isin(verdicts)].copy()
    if g.empty:
        return pd.DataFrame()
    out = (
        g.groupby(["target_id", "target_type", "target_title"], dropna=False)
        .agg(
            candidate_count=("pair_id", "count"),
            source_rule_count=("source_id", "nunique"),
        )
        .reset_index()
    )
    by_source = (
        g.groupby(["target_id", "target_type", "source"], dropna=False)
        .agg(source_rule_count=("source_id", "nunique"))
        .reset_index()
        .pivot_table(
            index=["target_id", "target_type"],
            columns="source",
            values="source_rule_count",
            fill_value=0,
        )
        .reset_index()
    )
    if "hadolint" not in by_source:
        by_source["hadolint"] = 0
    if "shellcheck" not in by_source:
        by_source["shellcheck"] = 0
    by_source = by_source.rename(columns={
        "hadolint": "hadolint_rule_count",
        "shellcheck": "shellcheck_rule_count",
    })
    out = out.merge(
        by_source[
            ["target_id", "target_type",
             "hadolint_rule_count", "shellcheck_rule_count"]
        ],
        on=["target_id", "target_type"],
        how="left",
    )
    return out.sort_values(
        ["candidate_count", "source_rule_count", "target_id"],
        ascending=[False, False, True],
    ).head(limit)


def make_analysis(df: pd.DataFrame, workbook_metadata: dict[str, Any]):
    total = len(df)
    rules = rule_summary(df)
    verdict_counts = df["llm_verdict"].value_counts(dropna=False)

    source_rows = []
    for source, g in df.groupby("source", sort=True):
        rg = rules[rules["source"] == source]
        source_rows.append({
            "source": source,
            "candidates": len(g),
            "candidate_percentage": pct(len(g), total),
            "rules": int(g["source_id"].nunique()),
            "yes": int((g["llm_verdict"] == "YES").sum()),
            "maybe": int((g["llm_verdict"] == "MAYBE").sum()),
            "no": int((g["llm_verdict"] == "NO").sum()),
            "yes_or_maybe": int(g["llm_verdict"].isin(["YES","MAYBE"]).sum()),
            "yes_percentage": pct((g["llm_verdict"] == "YES").sum(), len(g)),
            "maybe_percentage": pct((g["llm_verdict"] == "MAYBE").sum(), len(g)),
            "no_percentage": pct((g["llm_verdict"] == "NO").sum(), len(g)),
            "rules_with_yes": int(rg["has_yes"].sum()),
            "rules_with_yes_or_maybe": int(rg["has_yes_or_maybe"].sum()),
            "rules_all_no": int(rg["all_no"].sum()),
        })

    target_rows = []
    for target_type, g in df.groupby("target_type", sort=True):
        ym = g["llm_verdict"].isin(["YES", "MAYBE"]).sum()
        target_rows.append({
            "target_type": target_type,
            "candidates": len(g),
            "candidate_percentage": pct(len(g), total),
            "yes": int((g["llm_verdict"] == "YES").sum()),
            "maybe": int((g["llm_verdict"] == "MAYBE").sum()),
            "no": int((g["llm_verdict"] == "NO").sum()),
            "yes_or_maybe": int(ym),
            "yes_or_maybe_percentage": pct(ym, len(g)),
        })

    score_cols = {
        "rank": "rank",
        "raw_cosine": "raw_cosine",
        "relative_score": "relative_score",
        "top1_top2_gap": "top1_top2_gap",
        "llm_confidence": "llm_confidence",
    }
    score_stats = {}
    for name, col in score_cols.items():
        score_stats[name] = {
            "overall": stats(df[col]),
            "by_verdict": {
                v: stats(df.loc[df["llm_verdict"] == v, col])
                for v in VERDICTS
            },
        }

    integrity = {
        "expected_total": EXPECTED_TOTAL,
        "actual_total": total,
        "total_matches_expected": total == EXPECTED_TOTAL,
        "unique_pair_ids": int(df["pair_id"].nunique()),
        "duplicate_pair_ids": int(df["pair_id"].duplicated().sum()),
        "null_verdicts": int(df["llm_verdict"].isna().sum()),
        "unexpected_verdicts": sorted(
            set(df["llm_verdict"].dropna()) - set(VERDICTS)
        ),
        "source_counts": {
            str(k): int(v) for k, v in df["source"].value_counts().items()
        },
        "target_type_counts": {
            str(k): int(v) for k, v in df["target_type"].value_counts().items()
        },
        "source_counts_match_expected": (
            {str(k): int(v) for k, v in df["source"].value_counts().items()}
            == EXPECTED_SOURCE_COUNTS
        ),
        "target_type_counts_match_expected": (
            {str(k): int(v) for k, v in df["target_type"].value_counts().items()}
            == EXPECTED_TARGET_COUNTS
        ),
    }

    analysis = {
        "metadata": {
            "input_xlsx": str(INPUT_XLSX),
            "purpose": "Descriptive global analysis of the Gemini preliminary audit",
            "human_review_authoritative": True,
            "workbook_metadata": workbook_metadata,
        },
        "integrity": integrity,
        "overall": {
            "total_candidates": total,
            "unique_rules": int(df["source_id"].nunique()),
            "unique_targets": int(df["target_id"].nunique()),
            "human_reviewed_candidates": int(df["human_label"].notna().sum()),
            "verdicts": {
                v: int(verdict_counts.get(v, 0)) for v in VERDICTS
            },
            "verdict_distribution": distribution(df["llm_verdict"]),
            "yes_or_maybe": int(
                df["llm_verdict"].isin(["YES","MAYBE"]).sum()
            ),
            "yes_or_maybe_percentage": pct(
                df["llm_verdict"].isin(["YES","MAYBE"]).sum(), total
            ),
        },
        "by_source": source_rows,
        "by_target_type": target_rows,
        "cross_distributions": {
            "verdict_by_source": cross_distribution(df, "source", "llm_verdict"),
            "verdict_by_target_type": cross_distribution(
                df, "target_type", "llm_verdict"
            ),
            "verdict_by_rank": cross_distribution(
                df, "rank", "llm_verdict"
            ),
            "verdict_by_gap_bin": cross_distribution(
                df, "gap_bin", "llm_verdict"
            ),
            "verdict_by_rationale_availability": cross_distribution(
                df, "target_rationale_available", "llm_verdict"
            ),
        },
        "scores": score_stats,
        "directness": {
            "raw": distribution(df["llm_directness"]),
        },
        "relation_type": {
            "raw": distribution(df["llm_relation_type"]),
        },
        "security_objective": distribution(df["llm_security_objective"]),
        "rule_level": {
            "total_rules": len(rules),
            "rules_with_yes": int(rules["has_yes"].sum()),
            "rules_with_maybe": int(rules["has_maybe"].sum()),
            "rules_with_yes_or_maybe": int(rules["has_yes_or_maybe"].sum()),
            "rules_all_no": int(rules["all_no"].sum()),
            "rules_with_yes_percentage": pct(
                rules["has_yes"].sum(), len(rules)
            ),
            "rules_with_yes_or_maybe_percentage": pct(
                rules["has_yes_or_maybe"].sum(), len(rules)
            ),
        },
        "top_targets": {
            "yes": top_targets(df, {"YES"}).to_dict(orient="records"),
            "yes_or_maybe": top_targets(
                df, {"YES", "MAYBE"}
            ).to_dict(orient="records"),
        },
        "notes": [
            "Gemini verdicts are preliminary audit outputs.",
            "Human review remains authoritative for the gold standard.",
            "YES/MAYBE are not validated gold-standard labels.",
            "Rule-level coverage means that at least one candidate for a rule received the indicated Gemini verdict.",
            "Raw model-generated relation_type and directness values are reported without semantic reclassification.",
        ],
    }
    return analysis, rules


def excel_safe(v: Any):
    if v is None:
        return None
    if isinstance(v, (list, dict, tuple, set)):
        return json.dumps(v, ensure_ascii=False)
    if isinstance(v, float) and math.isnan(v):
        return None
    return v


def add_df_sheet(wb: Workbook, name: str, df: pd.DataFrame):
    ws = wb.create_sheet(name[:31])
    if df.empty:
        ws.append(["No data"])
        return
    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append([excel_safe(v) for v in row])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(df.columns, 1):
        width = min(
            60,
            max(
                10,
                len(str(col)) + 2,
                max(
                    [len(str(x)) for x in df.iloc[:200, i-1].tolist()
                     if x is not None] or [0]
                ) + 2,
            ),
        )
        ws.column_dimensions[get_column_letter(i)].width = width


def add_dict_sheet(wb: Workbook, name: str, d: dict[str, Any]):
    rows = []
    for k, v in d.items():
        if isinstance(v, dict):
            for sk, sv in v.items():
                rows.append({"metric": f"{k}.{sk}", "value": sv})
        else:
            rows.append({"metric": k, "value": v})
    add_df_sheet(wb, name, pd.DataFrame(rows))


def save_outputs(analysis, df, rules):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    OUTPUT_JSON.write_text(
        json.dumps(analysis, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    o = analysis["overall"]
    v = o["verdicts"]
    r = analysis["rule_level"]

    md = [
        "# Global Analysis — Gemini Preliminary Audit",
        "",
        "> Gemini is a preliminary auditor; human review remains authoritative.",
        "",
        "## Dataset",
        f"- Candidates: **{o['total_candidates']}**",
        f"- Unique rules: **{o['unique_rules']}**",
        f"- Unique IEC targets: **{o['unique_targets']}**",
        f"- Human-reviewed candidates: **{o['human_reviewed_candidates']}**",
        "",
        "## Gemini verdicts",
        "",
        "| Verdict | Count | Percentage |",
        "|---|---:|---:|",
    ]
    for vname in VERDICTS:
        md.append(
            f"| {vname} | {v[vname]} | "
            f"{pct(v[vname], o['total_candidates']):.2f}% |"
        )
    md += [
        f"| YES + MAYBE | {o['yes_or_maybe']} | "
        f"{o['yes_or_maybe_percentage']:.2f}% |",
        "",
        "## By source",
        "",
        "| Source | Candidates | YES | MAYBE | NO | YES/MAYBE | Rules | Rules with YES/MAYBE | Rules all NO |",
        "|---|---:|---:|---:|---:|---:|---:|---:|---:|",
    ]
    for x in analysis["by_source"]:
        md.append(
            f"| {x['source']} | {x['candidates']} | {x['yes']} | "
            f"{x['maybe']} | {x['no']} | {x['yes_or_maybe']} | "
            f"{x['rules']} | {x['rules_with_yes_or_maybe']} | "
            f"{x['rules_all_no']} |"
        )

    md += [
        "",
        "## By IEC target type",
        "",
        "| Type | Candidates | YES | MAYBE | NO | YES/MAYBE | YES/MAYBE within type |",
        "|---|---:|---:|---:|---:|---:|---:|",
    ]
    for x in analysis["by_target_type"]:
        md.append(
            f"| {x['target_type']} | {x['candidates']} | {x['yes']} | "
            f"{x['maybe']} | {x['no']} | {x['yes_or_maybe']} | "
            f"{x['yes_or_maybe_percentage']:.2f}% |"
        )

    md += [
        "",
        "## Confidence",
        "",
        f"- Overall: `{analysis['scores']['llm_confidence']['overall']}`",
    ]
    for vname in VERDICTS:
        md.append(
            f"- {vname}: `{analysis['scores']['llm_confidence']['by_verdict'][vname]}`"
        )

    md += [
        "",
        "## Rule-level coverage",
        f"- Total rules: **{r['total_rules']}**",
        f"- Rules with YES: **{r['rules_with_yes']}** "
        f"({r['rules_with_yes_percentage']:.2f}%)",
        f"- Rules with YES/MAYBE: **{r['rules_with_yes_or_maybe']}** "
        f"({r['rules_with_yes_or_maybe_percentage']:.2f}%)",
        f"- Rules all NO: **{r['rules_all_no']}**",
        "",
        "## Integrity",
        f"- Expected candidates: **{analysis['integrity']['expected_total']}**",
        f"- Actual candidates: **{analysis['integrity']['actual_total']}**",
        f"- Unique pair IDs: **{analysis['integrity']['unique_pair_ids']}**",
        f"- Duplicate pair IDs: **{analysis['integrity']['duplicate_pair_ids']}**",
        f"- Null verdicts: **{analysis['integrity']['null_verdicts']}**",
        f"- Source counts match expected: **{analysis['integrity']['source_counts_match_expected']}**",
        f"- Target-type counts match expected: **{analysis['integrity']['target_type_counts_match_expected']}**",
        "",
        "## Methodological notes",
        "- This is a descriptive analysis of Gemini's preliminary audit.",
        "- It does not establish the final gold standard.",
        "- No semantic relabeling of relation_type or directness was performed.",
        "- Rule-level YES/MAYBE coverage is not validated mapping coverage.",
        "",
    ]
    OUTPUT_MD.write_text("\n".join(md), encoding="utf-8")

    wb = Workbook()
    wb.remove(wb.active)

    overview = pd.DataFrame([
        ["Total candidates", o["total_candidates"]],
        ["YES", v["YES"]],
        ["MAYBE", v["MAYBE"]],
        ["NO", v["NO"]],
        ["YES + MAYBE", o["yes_or_maybe"]],
        ["YES + MAYBE (%)", o["yes_or_maybe_percentage"]],
        ["Unique rules", o["unique_rules"]],
        ["Unique targets", o["unique_targets"]],
        ["Human reviewed", o["human_reviewed_candidates"]],
        ["Rules with YES", r["rules_with_yes"]],
        ["Rules with YES/MAYBE", r["rules_with_yes_or_maybe"]],
        ["Rules all NO", r["rules_all_no"]],
    ], columns=["metric", "value"])
    add_df_sheet(wb, "Overview", overview)
    add_df_sheet(wb, "By Source", pd.DataFrame(analysis["by_source"]))
    add_df_sheet(wb, "By Target Type", pd.DataFrame(analysis["by_target_type"]))

    for key, name in [
        ("verdict_by_source", "Verdict x Source"),
        ("verdict_by_target_type", "Verdict x Target"),
        ("verdict_by_rank", "Verdict x Rank"),
        ("verdict_by_gap_bin", "Verdict x Gap"),
        ("verdict_by_rationale_availability", "Verdict x Rationale"),
    ]:
        add_df_sheet(
            wb, name,
            pd.DataFrame(analysis["cross_distributions"][key])
        )

    for metric, data in analysis["scores"].items():
        rows = [{"scope": "overall", **data["overall"]}]
        rows += [
            {"scope": vname, **data["by_verdict"][vname]}
            for vname in VERDICTS
        ]
        add_df_sheet(wb, metric[:31], pd.DataFrame(rows))

    add_df_sheet(wb, "Directness", pd.DataFrame(analysis["directness"]["raw"]))
    add_df_sheet(wb, "Relation Type", pd.DataFrame(analysis["relation_type"]["raw"]))
    add_df_sheet(wb, "Security Objectives", pd.DataFrame(analysis["security_objective"]))
    add_df_sheet(wb, "Rule Level", rules)
    add_df_sheet(wb, "Top YES Targets", pd.DataFrame(analysis["top_targets"]["yes"]))
    add_df_sheet(wb, "Top YES MAYBE Targets", pd.DataFrame(analysis["top_targets"]["yes_or_maybe"]))
    add_dict_sheet(wb, "Integrity", analysis["integrity"])
    add_dict_sheet(wb, "Metadata", analysis["metadata"])

    tmp = OUTPUT_XLSX.with_name(OUTPUT_XLSX.name + ".tmp")
    wb.save(tmp)
    tmp.replace(OUTPUT_XLSX)


def main():
    input_path = INPUT_XLSX
    if not input_path.exists() and FALLBACK_INPUT_XLSX.exists():
        input_path = FALLBACK_INPUT_XLSX
    if not input_path.exists():
        raise FileNotFoundError(f"Input XLSX not found: {INPUT_XLSX}")

    wb = load_workbook(input_path, read_only=True, data_only=True)
    if "Review" not in wb.sheetnames:
        raise ValueError("Required sheet 'Review' not found.")

    ws = wb["Review"]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Review sheet is empty.")

    columns = list(rows[0])
    df = pd.DataFrame(rows[1:], columns=columns)

    required = {
        "pair_id", "source", "source_id", "source_title", "rank",
        "target_id", "target_type", "target_title", "raw_cosine",
        "relative_score", "top1_top2_gap", "gap_bin",
        "target_rationale_available", "llm_verdict",
        "llm_relation_type", "llm_directness", "llm_confidence",
        "llm_security_objective", "human_label",
    }
    missing = sorted(required - set(columns))
    if missing:
        raise ValueError("Missing Review columns: " + ", ".join(missing))

    for col in ["rank", "raw_cosine", "relative_score", "top1_top2_gap",
                "llm_confidence"]:
        df[col] = pd.to_numeric(df[col], errors="coerce")

    for col in ["pair_id", "source", "source_id", "source_title",
                "target_id", "target_type", "target_title", "gap_bin",
                "target_rationale_available", "llm_verdict",
                "llm_relation_type", "llm_directness",
                "llm_security_objective", "human_label"]:
        df[col] = df[col].apply(clean)

    metadata = {}
    if "Metadata" in wb.sheetnames:
        mw = wb["Metadata"]
        for row in mw.iter_rows(values_only=True):
            if len(row) >= 2 and row[0] is not None:
                metadata[str(row[0])] = row[1]
    wb.close()

    analysis, rules = make_analysis(df, metadata)
    analysis["metadata"]["resolved_input_xlsx"] = str(input_path)

    bad = analysis["integrity"]["unexpected_verdicts"]
    if bad:
        raise ValueError("Unexpected verdicts: " + ", ".join(bad))
    if analysis["integrity"]["duplicate_pair_ids"]:
        raise ValueError("Duplicate pair_id values detected.")

    save_outputs(analysis, df, rules)

    print("=" * 60)
    print("GEMINI GLOBAL ANALYSIS")
    print("=" * 60)
    print(f"Candidates: {analysis['overall']['total_candidates']}")
    print(f"YES:       {analysis['overall']['verdicts']['YES']}")
    print(f"MAYBE:     {analysis['overall']['verdicts']['MAYBE']}")
    print(f"NO:        {analysis['overall']['verdicts']['NO']}")
    print(
        f"YES+MAYBE: {analysis['overall']['yes_or_maybe']} "
        f"({analysis['overall']['yes_or_maybe_percentage']:.2f}%)"
    )
    print(
        f"Rules with YES/MAYBE: "
        f"{analysis['rule_level']['rules_with_yes_or_maybe']}/"
        f"{analysis['rule_level']['total_rules']}"
    )
    print("\nOutputs:")
    print(f"  JSON: {OUTPUT_JSON}")
    print(f"  MD:   {OUTPUT_MD}")
    print(f"  XLSX: {OUTPUT_XLSX}")


if __name__ == "__main__":
    main()
